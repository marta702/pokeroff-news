#!/usr/bin/env python3
"""Сбор новостей по списку источников из sources.yaml.

  python news_parser.py                        # собрать всё: новые записи -> news.db + out/new_*.csv
  python news_parser.py --only telegram        # только один тип или группа (telegram, youtube, competitor, ...)
  python news_parser.py --source "PokerNews"   # один источник по имени
  python news_parser.py --check                # прогон без записи, только статус источников
  python news_parser.py --export news.xlsx --days 7   # выгрузка из базы в Excel за N дней
  python news_parser.py --html news.html --days 7     # лента новостей в браузере
"""
import argparse
import csv
import hashlib
import json
import logging
import os
import re
import sqlite3
import threading
from concurrent.futures import ThreadPoolExecutor, as_completed
from datetime import datetime, timedelta, timezone
from urllib.parse import quote_plus, urljoin, urlparse

import feedparser
import requests
import yaml
from bs4 import BeautifulSoup

BASE = os.path.dirname(os.path.abspath(__file__))
DB_PATH = os.path.join(BASE, "news.db")
OUT_DIR = os.path.join(BASE, "out")
FEED_CACHE = os.path.join(BASE, "feeds_cache.json")
TIMEOUT = 25
UA = ("Mozilla/5.0 (Windows NT 10.0; Win64; x64) AppleWebKit/537.36 "
      "(KHTML, like Gecko) Chrome/126.0 Safari/537.36")

log = logging.getLogger("news")
session = requests.Session()
session.headers.update({"User-Agent": UA, "Accept-Language": "ru,en;q=0.8"})


class SkipSource(Exception):
    """Источник пропущен из-за отсутствия ключей/настроек."""


# ---------- helpers ----------

def get(url, **kw):
    r = session.get(url, timeout=TIMEOUT, **kw)
    if is_challenge(r):
        raise RuntimeError("антибот-защита (Cloudflare): нужен браузер или другой канал")
    r.raise_for_status()
    return r


def is_challenge(r):
    if r.status_code not in (403, 429, 503):
        return False
    t = r.text[:5000]
    return any(m in t for m in ("cf-chl", "challenge-platform", "Just a moment", "cf_chl_opt"))


def clean(text, limit=1000):
    return re.sub(r"\s+", " ", text or "").strip()[:limit]


def html_to_text(html):
    return clean(BeautifulSoup(html or "", "lxml").get_text(" "))


def to_iso(value):
    if not value:
        return None
    if isinstance(value, datetime):
        dt = value
    else:
        try:
            dt = datetime.fromisoformat(str(value).replace("Z", "+00:00"))
        except ValueError:
            return str(value)
    if dt.tzinfo is None:
        dt = dt.replace(tzinfo=timezone.utc)
    return dt.astimezone(timezone.utc).isoformat(timespec="seconds")


def item(src, title, url, published=None, text="", author=""):
    return {
        "source": src["name"], "group": src.get("group", ""), "type": src["type"],
        "title": clean(title, 300), "url": url, "published": to_iso(published),
        "author": clean(author, 120), "text": clean(text),
    }


def load_feed_cache():
    try:
        with open(FEED_CACHE, encoding="utf-8") as f:
            return json.load(f)
    except (OSError, ValueError):
        return {}


_feed_cache = load_feed_cache()
_cache_lock = threading.Lock()


# ---------- fetchers ----------

def parse_feed(src, content):
    feed = feedparser.parse(content)
    if not feed.entries:
        raise RuntimeError("лента пустая или это не RSS/Atom")
    out = []
    for e in feed.entries:
        t = e.get("published_parsed") or e.get("updated_parsed")
        pub = datetime(*t[:6], tzinfo=timezone.utc) if t else None
        body = e.get("summary", "")
        if e.get("content"):
            body = e.content[0].get("value", body)
        author = e.get("author", "") or (e.get("source") or {}).get("title", "")
        out.append(item(src, e.get("title", ""), e.get("link", ""), pub, html_to_text(body), author))
    return out


def fetch_rss(src):
    return parse_feed(src, get(src["url"]).content)


def feed_candidates(page_url):
    p = urlparse(page_url)
    root = f"{p.scheme}://{p.netloc}/"
    page_dir = page_url if page_url.endswith("/") else page_url.rsplit("/", 1)[0] + "/"
    names = ["feed/", "rss/", "rss.xml", "feed.xml", "atom.xml", "index.xml", "rss"]
    seen, out = set(), []
    for b in (page_dir, root):
        for n in names:
            u = urljoin(b, n)
            if u not in seen:
                seen.add(u)
                out.append(u)
    return out


def discover_feed(url, resp):
    ctype = resp.headers.get("content-type", "")
    if "xml" in ctype or resp.text.lstrip().startswith("<?xml"):
        return url
    soup = BeautifulSoup(resp.text, "lxml")
    for link in soup.find_all("link", href=True):
        rel = " ".join(link.get("rel") or []).lower()
        typ = (link.get("type") or "").lower()
        if "alternate" in rel and ("rss" in typ or "atom" in typ):
            return urljoin(resp.url, link["href"])
    for cand in feed_candidates(resp.url):
        try:
            r = session.get(cand, timeout=TIMEOUT)
            if r.ok and feedparser.parse(r.content).entries:
                return cand
        except requests.RequestException:
            pass
    return None


def fetch_auto(src):
    """Сначала RSS (автопоиск, результат кэшируется), иначе разбор HTML."""
    url = src["url"]
    cached = _feed_cache.get(url)
    if cached:
        try:
            return parse_feed(src, get(cached).content)
        except Exception as e:  # лента сломалась — ищем заново
            log.debug("кэш ленты %s не сработал: %s", cached, e)
    resp = get(url)
    feed_url = discover_feed(url, resp)
    with _cache_lock:
        _feed_cache[url] = feed_url
    if feed_url:
        try:
            items = parse_feed(src, get(feed_url).content)
            src["_via"] = f"rss: {feed_url}"
            return items
        except Exception as e:  # сайт указывает нерабочую ленту — переходим к HTML
            log.debug("лента %s не работает: %s", feed_url, e)
            with _cache_lock:
                _feed_cache[url] = None
    src["_via"] = "html"
    return fetch_html(src, resp)


ARTICLE_PATH = re.compile(r"\d{3,}|/[a-z0-9]+(?:-[a-z0-9]+){2,}", re.I)


def fetch_html(src, resp=None):
    """Эвристика: ссылки того же домена, похожие на статьи/темы, с заголовком от N символов.
    Точнее — задать в sources.yaml selector (CSS для <a>) и/или url_pattern (regex)."""
    resp = resp or get(src["url"])
    soup = BeautifulSoup(resp.text, "lxml")
    host = urlparse(resp.url).netloc.removeprefix("www.")
    pattern = re.compile(src["url_pattern"]) if src.get("url_pattern") else None
    min_len = src.get("min_title", 25)
    nodes = soup.select(src["selector"]) if src.get("selector") else soup.find_all("a", href=True)
    seen, out = set(), []
    for a in nodes:
        if a.name != "a":
            a = a.find("a", href=True)
            if a is None:
                continue
        href = a.get("href")
        if not href:
            continue
        url = urljoin(resp.url, href).split("#")[0]
        p = urlparse(url)
        if p.netloc.removeprefix("www.") != host:
            continue
        if pattern:
            if not pattern.search(url):
                continue
        elif not ARTICLE_PATH.search(p.path):
            continue
        title = clean(a.get("title") or a.get_text(" "), 300)
        if len(title) < min_len or url in seen:
            continue
        seen.add(url)
        out.append(item(src, title, url))
    if not out:
        raise RuntimeError("ничего не найдено: нужен selector/url_pattern в sources.yaml")
    return out[: src.get("limit", 50)]


def fetch_telegram(src):
    r = get(f"https://t.me/s/{src['channel']}")
    soup = BeautifulSoup(r.text, "lxml")
    msgs = soup.select("div.tgme_widget_message[data-post]")
    if not msgs:
        raise RuntimeError("нет постов: канал закрыт, не существует или веб-превью отключено")
    out = []
    for m in msgs:
        el = m.select_one(".tgme_widget_message_text")
        text = el.get_text("\n") if el else ""
        t = m.select_one("time[datetime]")
        first_line = next((ln for ln in text.split("\n") if ln.strip()), "")
        title = clean(first_line, 150) or "[медиа без текста]"
        out.append(item(src, title, f"https://t.me/{m['data-post']}", t["datetime"] if t else None, text))
    return out


def fetch_youtube(src):
    if src.get("channel_id"):
        url = f"https://www.youtube.com/feeds/videos.xml?channel_id={src['channel_id']}"
    else:
        url = f"https://www.youtube.com/feeds/videos.xml?user={src['user']}"
    return parse_feed(src, get(url).content)


def fetch_reddit(src):
    return parse_feed(src, get(f"https://www.reddit.com/r/{src['subreddit']}/new/.rss").content)


def fetch_google_news(src):
    url = (f"https://news.google.com/rss/search?q={quote_plus(src['query'])}"
           f"&hl={src.get('hl', 'ru')}&gl={src.get('gl', 'RU')}&ceid={quote_plus(src.get('ceid', 'RU:ru'))}")
    return parse_feed(src, get(url).content)


_tw = {"token": None}
_tw_lock = threading.Lock()


def twitch_headers():
    cid, sec = os.getenv("TWITCH_CLIENT_ID"), os.getenv("TWITCH_CLIENT_SECRET")
    if not (cid and sec):
        raise SkipSource("нет TWITCH_CLIENT_ID / TWITCH_CLIENT_SECRET")
    with _tw_lock:
        if not _tw["token"]:
            r = requests.post("https://id.twitch.tv/oauth2/token", timeout=TIMEOUT, data={
                "client_id": cid, "client_secret": sec, "grant_type": "client_credentials"})
            r.raise_for_status()
            _tw["token"] = r.json()["access_token"]
    return {"Client-ID": cid, "Authorization": f"Bearer {_tw['token']}"}


def helix(path, params):
    r = requests.get(f"https://api.twitch.tv/helix/{path}", headers=twitch_headers(),
                     params=params, timeout=TIMEOUT)
    r.raise_for_status()
    return r.json()["data"]


def fetch_twitch(src):
    if src.get("game"):
        games = helix("games", {"name": src["game"]})
        if not games:
            raise RuntimeError("категория не найдена")
        since = datetime.now(timezone.utc) - timedelta(days=src.get("days", 7))
        clips = helix("clips", {"game_id": games[0]["id"], "started_at": to_iso(since),
                                "first": src.get("limit", 50)})
        return [item(src, c["title"], c["url"], c["created_at"],
                     f"{c['view_count']} просмотров", c["broadcaster_name"]) for c in clips]
    users = helix("users", {"login": src["channel"]})
    if not users:
        raise RuntimeError("канал не найден")
    videos = helix("videos", {"user_id": users[0]["id"], "first": src.get("limit", 20)})
    return [item(src, v["title"], v["url"], v["created_at"], v.get("description", ""), v["user_name"])
            for v in videos]


def fetch_rsshub(src):
    """X и Instagram: через свой инстанс RSSHub (RSSHUB_URL). Либо замените тип на rss
    и укажите url ленты из RSS.app / аналога."""
    base = os.getenv("RSSHUB_URL")
    if not base:
        raise SkipSource("не задан RSSHUB_URL")
    return parse_feed(src, get(base.rstrip("/") + src["route"]).content)


FETCHERS = {
    "rss": fetch_rss, "auto": fetch_auto, "html": fetch_html, "telegram": fetch_telegram,
    "youtube": fetch_youtube, "reddit": fetch_reddit, "google_news": fetch_google_news,
    "twitch": fetch_twitch, "rsshub": fetch_rsshub,
}


def apply_keywords(src, items):
    kws = [k.lower() for k in src.get("keywords") or []]
    if not kws:
        return items
    return [i for i in items if any(k in (i["title"] + " " + i["text"]).lower() for k in kws)]


# ---------- storage ----------

COLUMNS = ["source", "group", "type", "title", "url", "published", "author", "text", "fetched_at"]


def db_connect():
    conn = sqlite3.connect(DB_PATH)
    conn.execute("""CREATE TABLE IF NOT EXISTS items (
        id TEXT PRIMARY KEY, source TEXT, "group" TEXT, type TEXT, title TEXT, url TEXT,
        published TEXT, author TEXT, text TEXT, fetched_at TEXT)""")
    return conn


def item_id(i):
    key = i["url"] or f"{i['source']}|{i['title']}"
    return hashlib.sha1(key.encode("utf-8")).hexdigest()


def save_new(conn, items):
    now = to_iso(datetime.now(timezone.utc))
    new = []
    for i in items:
        i["fetched_at"] = now
        cur = conn.execute(
            'INSERT OR IGNORE INTO items (id, source, "group", type, title, url, published, author, text, fetched_at) '
            "VALUES (?,?,?,?,?,?,?,?,?,?)",
            (item_id(i), *[i[c] for c in COLUMNS]))
        if cur.rowcount:
            new.append(i)
    conn.commit()
    return new


def write_csv(path, rows, columns):
    os.makedirs(os.path.dirname(path), exist_ok=True)
    with open(path, "w", newline="", encoding="utf-8-sig") as f:  # utf-8-sig — корректно открывается в Excel
        w = csv.DictWriter(f, fieldnames=columns, extrasaction="ignore")
        w.writeheader()
        w.writerows(rows)


def export_xlsx(path, days):
    from openpyxl import Workbook
    from openpyxl.styles import Font
    since = to_iso(datetime.now(timezone.utc) - timedelta(days=days))
    conn = db_connect()
    rows = conn.execute(
        f'SELECT {", ".join(chr(34) + c + chr(34) for c in COLUMNS)} FROM items '
        "WHERE COALESCE(published, fetched_at) >= ? ORDER BY COALESCE(published, fetched_at) DESC",
        (since,)).fetchall()
    wb = Workbook()
    ws = wb.active
    ws.title = "Новости"
    ws.append(COLUMNS)
    for c in ws[1]:
        c.font = Font(bold=True)
    for r in rows:
        ws.append(list(r))
        cell = ws.cell(row=ws.max_row, column=COLUMNS.index("url") + 1)
        if cell.value:
            cell.hyperlink = cell.value
            cell.font = Font(color="0563C1", underline="single")
    for col, width in zip("ABCDEFGHI", (22, 12, 10, 70, 45, 22, 20, 80, 22)):
        ws.column_dimensions[col].width = width
    ws.auto_filter.ref = ws.dimensions
    ws.freeze_panes = "A2"
    wb.save(path)
    print(f"{len(rows)} записей за {days} дн. -> {path}")


def export_html(path, days):
    since = to_iso(datetime.now(timezone.utc) - timedelta(days=days))
    conn = db_connect()
    rows = conn.execute(
        'SELECT source, "group", type, title, url, COALESCE(published, fetched_at), author, text FROM items '
        "WHERE COALESCE(published, fetched_at) >= ? ORDER BY COALESCE(published, fetched_at) DESC",
        (since,)).fetchall()
    data = [dict(zip(("s", "g", "t", "ti", "u", "p", "a", "x"), r)) for r in rows]
    payload = json.dumps(data, ensure_ascii=False).replace("</", "<\\/")
    updated = to_iso(datetime.now(timezone.utc))
    html = (HTML_TEMPLATE.replace("__DATA__", payload)
            .replace("__DAYS__", str(days)).replace("__UPDATED__", updated))
    if os.path.dirname(path):
        os.makedirs(os.path.dirname(path), exist_ok=True)
    with open(path, "w", encoding="utf-8") as f:
        f.write(html)
    print(f"{len(data)} записей за {days} дн. -> {path}")


HTML_TEMPLATE = r"""<!doctype html>
<html lang="ru">
<head>
<meta charset="utf-8">
<meta name="viewport" content="width=device-width, initial-scale=1">
<title>Покерные новости</title>
<style>
:root{
  --felt:#1d4a3a; --felt-2:#153a2d; --paper:#f5f6f3; --card:#ffffff; --ink:#1c2420; --muted:#66706b;
  --line:#dfe3dd; --focus:#c9a227;
  --c-competitor:#c0392b; --c-room:#2e5fa8; --c-media:#8a8375; --c-community:#2f8a5b;
  --c-influencer:#7b4fa6; --c-own:#c9a227;
}
@media (prefers-color-scheme:dark){
  :root{--paper:#151a18; --card:#1d2421; --ink:#e7ebe8; --muted:#95a09a; --line:#2c3531; --felt:#17392d; --felt-2:#102a21}
}
*{box-sizing:border-box}
body{margin:0;background:var(--paper);color:var(--ink);font:15px/1.5 -apple-system,BlinkMacSystemFont,"Segoe UI",Roboto,sans-serif}
header{background:var(--felt);color:#f3efe2;padding:22px 28px 18px;border-bottom:4px solid var(--felt-2)}
header h1{margin:0;font:600 26px/1.2 "Iowan Old Style","Charter",Georgia,serif;letter-spacing:.2px}
header p{margin:4px 0 0;color:#bcd0c6;font-size:13px}
.wrap{display:grid;grid-template-columns:250px minmax(0,1fr);gap:28px;max-width:1180px;margin:0 auto;padding:24px 28px 100px}
aside{position:sticky;top:16px;align-self:start}
aside h2{font-size:13px;font-weight:600;color:var(--muted);margin:22px 0 8px}
aside h2:first-child{margin-top:0}
input[type=search],select{width:100%;padding:9px 11px;border:1px solid var(--line);border-radius:8px;background:var(--card);color:var(--ink);font:inherit}
.chips{display:flex;flex-direction:column;gap:2px}
.chip{display:flex;align-items:center;gap:9px;width:100%;padding:6px 8px;border:0;border-radius:6px;background:none;color:var(--ink);font:inherit;text-align:left;cursor:pointer}
.chip:hover{background:var(--line)}
.chip .dot{width:12px;height:12px;border-radius:50%;flex:none;box-shadow:inset 0 0 0 2px rgba(255,255,255,.5)}
.chip .n{margin-left:auto;color:var(--muted);font-size:13px;font-variant-numeric:tabular-nums}
.chip[aria-pressed=false]{opacity:.4}
.period{display:flex;gap:4px}
.period button{flex:1;padding:7px 0;border:1px solid var(--line);border-radius:6px;background:var(--card);color:var(--ink);font:inherit;font-size:13px;cursor:pointer}
.period button[aria-pressed=true]{background:var(--felt);border-color:var(--felt);color:#fff}
button:focus-visible,input:focus-visible,select:focus-visible,a:focus-visible{outline:2px solid var(--focus);outline-offset:2px}
.reset{margin-top:14px;background:none;border:0;color:var(--muted);text-decoration:underline;cursor:pointer;font:inherit;font-size:13px;padding:0}
main{min-width:0}
.summary{color:var(--muted);font-size:13px;margin:0 0 6px}
.day{font:600 17px/1.3 "Iowan Old Style","Charter",Georgia,serif;margin:26px 0 10px;padding-bottom:6px;border-bottom:1px solid var(--line)}
.day:first-of-type{margin-top:8px}
article{background:var(--card);border:1px solid var(--line);border-left:4px solid var(--gc);border-radius:4px 8px 8px 4px;padding:12px 16px;margin-bottom:8px}
article .meta{display:flex;flex-wrap:wrap;gap:4px 12px;font-size:12.5px;color:var(--muted)}
article .src{color:var(--gc);font-weight:600}
article h3{margin:4px 0 0;font:600 16.5px/1.35 "Iowan Old Style","Charter",Georgia,serif;max-width:75ch}
article h3 a{color:inherit;text-decoration:none}
article h3 a:hover{text-decoration:underline}
article .x{margin:6px 0 0;color:var(--ink);opacity:.85;font-size:14px;max-width:75ch;display:-webkit-box;-webkit-line-clamp:2;-webkit-box-orient:vertical;overflow:hidden}
article.open .x{display:block}
article .more{margin-top:4px;background:none;border:0;padding:0;color:var(--muted);font:inherit;font-size:12.5px;cursor:pointer;text-decoration:underline}
.empty{padding:40px 0;color:var(--muted)}
.actions{display:flex;align-items:center;gap:14px;margin-top:8px}
.copy{padding:5px 11px;border:1px solid var(--felt);border-radius:6px;background:none;color:var(--felt);font:inherit;font-size:13px;cursor:pointer}
.copy:hover{background:var(--felt);color:#fff}
@media (prefers-color-scheme:dark){.copy{border-color:#6fae92;color:#9fd3bb}}
.pick{display:flex;align-items:center;gap:6px;font-size:13px;color:var(--muted);cursor:pointer}
.pick input{width:16px;height:16px;accent-color:var(--felt)}
article.picked{outline:2px solid var(--felt);outline-offset:-1px}
.bar{position:fixed;left:50%;bottom:18px;transform:translateX(-50%);display:none;align-items:center;gap:14px;background:var(--felt);color:#fff;padding:10px 12px 10px 18px;border-radius:10px;box-shadow:0 6px 24px rgba(0,0,0,.25);z-index:5}
.bar.show{display:flex}
.bar button{padding:7px 13px;border:0;border-radius:6px;font:inherit;font-size:13.5px;cursor:pointer}
.bar .go{background:#f3efe2;color:var(--felt-2);font-weight:600}
.bar .clear{background:transparent;color:#cfe0d7;text-decoration:underline;padding:7px 4px}
.toast{position:fixed;left:50%;top:18px;transform:translateX(-50%);background:var(--ink);color:var(--paper);padding:9px 16px;border-radius:8px;font-size:14px;opacity:0;transition:opacity .2s;pointer-events:none;z-index:6}
.toast.show{opacity:1}
@media (prefers-reduced-motion:reduce){.toast{transition:none}}
mark{background:#f3e3a3;color:inherit;border-radius:2px}
@media (max-width:820px){.wrap{grid-template-columns:1fr;padding:16px}aside{position:static}.chips{flex-direction:row;flex-wrap:wrap}.chip{width:auto}}
</style>
</head>
<body>
<header>
  <h1>Покерные новости</h1>
  <p>Обновлено <span id="upd" data-t="__UPDATED__"></span>, в ленте новости за последние __DAYS__ дн.</p>
</header>
<div class="wrap">
  <aside>
    <h2>Поиск</h2>
    <input type="search" id="q" placeholder="Слово в заголовке или тексте">
    <h2>Период</h2>
    <div class="period" id="period"></div>
    <h2>Группы</h2>
    <div class="chips" id="groups"></div>
    <h2>Формат</h2>
    <div class="chips" id="types"></div>
    <h2>Источник</h2>
    <select id="source"><option value="">Все источники</option></select>
    <button class="reset" id="reset">Сбросить фильтры</button>
  </aside>
  <main>
    <p class="summary" id="summary"></p>
    <div id="feed"></div>
  </main>
</div>
<div class="bar" id="bar"><span id="barN"></span><button class="go" id="copySel" type="button">Скопировать выбранные</button><button class="clear" id="clearSel" type="button">Снять выбор</button></div>
<div class="toast" id="toast" role="status"></div>
<script>
const DATA = __DATA__;
(u => { u.textContent = new Date(u.dataset.t).toLocaleString("ru-RU", {day:"numeric", month:"long", hour:"2-digit", minute:"2-digit"}); })(document.getElementById("upd"));
const GROUPS = {competitor:"Конкуренты", room:"Румы", media:"Медиа", community:"Сообщество", influencer:"Блогеры", own:"Pokeroff", "":"Без группы"};
const TYPES = {site:"Сайты", telegram:"Telegram", youtube:"YouTube", reddit:"Reddit", twitch:"Twitch", google_news:"Google News", rsshub:"X и Instagram"};
const typeOf = t => (t==="auto"||t==="rss"||t==="html") ? "site" : t;
const color = g => getComputedStyle(document.documentElement).getPropertyValue("--c-"+(g||"media")).trim() || "#888";
const state = {q:"", period:0, groups:new Set(), types:new Set(), source:""};
const PERIODS = [[1,"Сутки"],[3,"3 дня"],[0,"Все"]];

DATA.forEach((d, i) => { d.i = i; d.tt = typeOf(d.t); d.ts = Date.parse(d.p) || 0; d.hay = (d.ti+" "+(d.x||"")+" "+d.s).toLowerCase(); });
const esc = s => (s||"").replace(/[&<>"]/g, c => ({"&":"&amp;","<":"&lt;",">":"&gt;",'"':"&quot;"}[c]));
function hl(s){ const t = esc(s); if(!state.q) return t;
  const q = state.q.replace(/[.*+?^${}()|[\]\\]/g,"\\$&"); return t.replace(new RegExp(q,"gi"), m=>"<mark>"+m+"</mark>"); }

function passes(d, skip){
  if(state.q && !d.hay.includes(state.q)) return false;
  if(state.period && d.ts < Date.now() - state.period*864e5) return false;
  if(skip!=="g" && state.groups.size && !state.groups.has(d.g||"")) return false;
  if(skip!=="t" && state.types.size && !state.types.has(d.tt)) return false;
  if(state.source && d.s !== state.source) return false;
  return true;
}

function buildChips(el, dict, key, set, withDot){
  const present = [...new Set(DATA.map(d => key==="g" ? (d.g||"") : d.tt))];
  el.innerHTML = "";
  Object.keys(dict).filter(k => present.includes(k)).forEach(k => {
    const b = document.createElement("button");
    b.className = "chip"; b.dataset.k = k; b.setAttribute("aria-pressed","true");
    b.innerHTML = (withDot ? `<span class="dot" style="background:${color(k)}"></span>` : "") + `<span>${dict[k]}</span><span class="n"></span>`;
    b.onclick = () => {
      if(!set.size){ set.add(k); }            // первый клик: оставить только эту группу
      else if(set.has(k)){ set.delete(k); }
      else { set.add(k); }
      render();
    };
    el.appendChild(b);
  });
}

function dayLabel(ts){
  if(!ts) return "Без даты";
  const d = new Date(ts), t = new Date(); t.setHours(0,0,0,0);
  const diff = Math.round((t - new Date(d.getFullYear(), d.getMonth(), d.getDate())) / 864e5);
  if(diff === 0) return "Сегодня"; if(diff === 1) return "Вчера";
  return d.toLocaleDateString("ru-RU", {day:"numeric", month:"long", weekday:"long"});
}

function render(){
  for(const [el,key,set] of [[groupsEl,"g",state.groups],[typesEl,"t",state.types]]){
    el.querySelectorAll(".chip").forEach(b => {
      const k = b.dataset.k;
      b.querySelector(".n").textContent = DATA.filter(d => passes(d,key) && (key==="g" ? (d.g||"") : d.tt) === k).length;
      b.setAttribute("aria-pressed", !set.size || set.has(k) ? "true" : "false");
    });
  }
  const items = DATA.filter(d => passes(d));
  document.getElementById("summary").textContent = items.length ? `Показано ${items.length} из ${DATA.length}` : "";
  const feed = document.getElementById("feed");
  if(!items.length){ feed.innerHTML = `<p class="empty">Под эти фильтры новостей нет. Уберите часть фильтров или нажмите «Сбросить фильтры».</p>`; return; }
  let html = "", last = null;
  for(const d of items.slice(0, 600)){
    const lbl = dayLabel(d.ts);
    if(lbl !== last){ html += `<h2 class="day">${lbl}</h2>`; last = lbl; }
    const time = d.ts ? new Date(d.ts).toLocaleTimeString("ru-RU",{hour:"2-digit",minute:"2-digit"}) : "";
    const long = (d.x||"").length > 180;
    html += `<article class="${picked.has(d.i)?"picked":""}" style="--gc:${color(d.g)}">
      <div class="meta"><span class="src">${esc(d.s)}</span><span>${GROUPS[d.g||""]||""}</span><span>${TYPES[d.tt]||""}</span>${d.a && d.a!==d.s ? `<span>${esc(d.a)}</span>`:""}<span>${time}</span></div>
      <h3><a href="${esc(d.u)}" target="_blank" rel="noopener">${hl(d.ti)}</a></h3>
      ${d.x && d.x!==d.ti ? `<p class="x">${hl(d.x)}</p>` : ""}
      ${long ? `<button class="more" type="button">Показать полностью</button>` : ""}
      <div class="actions">
        <button class="copy" type="button" data-i="${d.i}">Скопировать для Claude</button>
        <label class="pick"><input type="checkbox" data-i="${d.i}" ${picked.has(d.i)?"checked":""}> Выбрать</label>
      </div>
    </article>`;
  }
  feed.innerHTML = html;
  feed.querySelectorAll(".copy").forEach(b => b.onclick = () =>
    copyText(prompt1([DATA[+b.dataset.i]]), "Скопировано. Вставьте в чат с Claude"));
  feed.querySelectorAll(".pick input").forEach(c => c.onchange = () => {
    const i = +c.dataset.i; c.checked ? picked.add(i) : picked.delete(i);
    c.closest("article").classList.toggle("picked", c.checked); updateBar();
  });
  feed.querySelectorAll(".more").forEach(b => b.onclick = () => {
    const a = b.closest("article"); a.classList.toggle("open");
    b.textContent = a.classList.contains("open") ? "Свернуть" : "Показать полностью";
  });
}

const picked = new Set();
function prompt1(list){
  const head = list.length > 1 ? `Перепиши эти ${list.length} новости в посты для Threads (отдельный пост на каждую):` : "Перепиши эту новость в пост для Threads:";
  return head + "\n\n" + list.map((d, k) => [
    list.length > 1 ? `[${k+1}] ${d.ti}` : d.ti,
    `Источник: ${d.s}${GROUPS[d.g||""] && d.g ? " (" + GROUPS[d.g] + ")" : ""}`,
    d.p ? `Дата: ${new Date(d.ts).toLocaleString("ru-RU")}` : "",
    d.x && d.x !== d.ti ? `Текст: ${d.x}` : "",
    `Ссылка: ${d.u}`
  ].filter(Boolean).join("\n")).join("\n\n");
}
function toast(msg){ const t = document.getElementById("toast"); t.textContent = msg; t.classList.add("show");
  clearTimeout(toast.t); toast.t = setTimeout(() => t.classList.remove("show"), 2200); }
async function copyText(text, okMsg){
  try { await navigator.clipboard.writeText(text); }
  catch(e){ const ta = document.createElement("textarea"); ta.value = text; document.body.appendChild(ta); ta.select();
    document.execCommand("copy"); ta.remove(); }
  toast(okMsg);
}
function updateBar(){
  document.getElementById("bar").classList.toggle("show", picked.size > 0);
  document.getElementById("barN").textContent = `Выбрано: ${picked.size}`;
}
document.getElementById("copySel").onclick = () => {
  const list = [...picked].sort((a,b) => a-b).map(i => DATA[i]);
  copyText(prompt1(list), `Скопировано ${list.length}. Вставьте в чат с Claude`);
};
document.getElementById("clearSel").onclick = () => { picked.clear(); updateBar(); render(); };

const groupsEl = document.getElementById("groups"), typesEl = document.getElementById("types");
buildChips(groupsEl, GROUPS, "g", state.groups, true);
buildChips(typesEl, TYPES, "t", state.types, false);
const periodEl = document.getElementById("period");
PERIODS.forEach(([v,l]) => { const b = document.createElement("button"); b.textContent = l; b.dataset.v = v;
  b.onclick = () => { state.period = v; periodEl.querySelectorAll("button").forEach(x => x.setAttribute("aria-pressed", x===b)); render(); };
  b.setAttribute("aria-pressed", v===0); periodEl.appendChild(b); });
const sel = document.getElementById("source");
[...new Set(DATA.map(d => d.s))].sort().forEach(s => sel.add(new Option(s, s)));
sel.onchange = () => { state.source = sel.value; render(); };
let tmr; document.getElementById("q").oninput = e => { clearTimeout(tmr); tmr = setTimeout(() => { state.q = e.target.value.trim().toLowerCase(); render(); }, 150); };
document.getElementById("reset").onclick = () => {
  state.q=""; state.period=0; state.groups.clear(); state.types.clear(); state.source="";
  document.getElementById("q").value=""; sel.value="";
  periodEl.querySelectorAll("button").forEach(x => x.setAttribute("aria-pressed", x.dataset.v==="0")); render(); };
render();
</script>
</body>
</html>
"""

# ---------- run ----------

def load_sources(path):
    with open(path, encoding="utf-8") as f:
        data = yaml.safe_load(f)
    seen, out = set(), []
    for s in data["sources"]:
        if s.get("disabled"):
            continue
        key = (s["type"], s.get("url") or s.get("channel") or s.get("channel_id") or s.get("user")
               or s.get("route") or s.get("query") or s.get("subreddit") or s.get("game"))
        if key in seen:
            continue
        seen.add(key)
        out.append(s)
    return out


def run_one(src):
    try:
        items = apply_keywords(src, FETCHERS[src["type"]](src))
        return src, items, "ok"
    except SkipSource as e:
        return src, [], f"пропущен: {e}"
    except Exception as e:
        return src, [], f"ошибка: {type(e).__name__}: {str(e)[:200]}"


def main():
    ap = argparse.ArgumentParser()
    ap.add_argument("--sources", default=os.path.join(BASE, "sources.yaml"))
    ap.add_argument("--only", help="тип или группа источников")
    ap.add_argument("--source", help="имя источника")
    ap.add_argument("--check", action="store_true", help="без записи в базу")
    ap.add_argument("--export", help="путь к .xlsx для выгрузки из базы")
    ap.add_argument("--html", help="путь к .html для ленты новостей из базы")
    ap.add_argument("--days", type=int, default=7)
    ap.add_argument("--workers", type=int, default=8)
    ap.add_argument("-v", action="store_true")
    a = ap.parse_args()
    logging.basicConfig(level=logging.DEBUG if a.v else logging.INFO, format="%(message)s")

    if a.export or a.html:
        if a.export:
            export_xlsx(a.export, a.days)
        if a.html:
            export_html(a.html, a.days)
        return

    sources = load_sources(a.sources)
    if a.only:
        sources = [s for s in sources if a.only in (s["type"], s.get("group"))]
    if a.source:
        sources = [s for s in sources if s["name"].lower() == a.source.lower()]

    conn = None if a.check else db_connect()
    status, all_new = [], []
    with ThreadPoolExecutor(max_workers=a.workers) as ex:
        for fut in as_completed([ex.submit(run_one, s) for s in sources]):
            src, items, st = fut.result()
            new = save_new(conn, items) if conn and items else []
            all_new += new
            status.append({"source": src["name"], "type": src["type"], "status": st,
                           "via": src.get("_via", ""), "found": len(items), "new": len(new)})
            log.info("%-28s %-11s найдено %-4d новых %-4d %s", src["name"][:28], src["type"],
                     len(items), len(new), "" if st == "ok" else st)

    with open(FEED_CACHE, "w", encoding="utf-8") as f:
        json.dump(_feed_cache, f, ensure_ascii=False, indent=1)
    stamp = datetime.now().strftime("%Y%m%d_%H%M")
    write_csv(os.path.join(OUT_DIR, "last_run_status.csv"), status,
              ["source", "type", "status", "via", "found", "new"])
    if all_new:
        all_new.sort(key=lambda i: i["published"] or i["fetched_at"], reverse=True)
        path = os.path.join(OUT_DIR, f"new_{stamp}.csv")
        write_csv(path, all_new, COLUMNS)
        log.info("\nНовых записей: %d -> %s", len(all_new), path)
    else:
        log.info("\nНовых записей нет")


if __name__ == "__main__":
    main()
